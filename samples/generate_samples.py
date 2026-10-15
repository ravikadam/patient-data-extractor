"""Generate a fictional, deliberately messy multi-format case file for one patient.

Patient: Eleanor M. Vasquez (fictional). The documents disagree on purpose:
  * metformin dose (500 mg in discharge summary -> 1000 mg in PCP e-mail / intake form)
  * furosemide dose (40 mg -> 20 mg on the cardiology prescription; intake form says 20 mg)
  * mobile phone (old in discharge summary, new in e-mail/intake/HL7)
  * address (patient moved; discharge summary has the old one)
  * penicillin reaction (hives vs. lip swelling)
  * date of birth typo on the patient-filled intake form (06/23 vs 06/22)
  * lisinopril discontinued and replaced by sacubitril/valsartan
Also contains material for custom-field demos (preferred pharmacy, code status / advance directive).

Run:  .venv/bin/python samples/generate_samples.py
"""
from __future__ import annotations

import csv
import io
import random
from email.message import EmailMessage
from pathlib import Path

OUT = Path(__file__).resolve().parent / "case_eleanor_vasquez"
FONTS = "/System/Library/Fonts"


def font(name: str, size: int):
    from PIL import ImageFont

    for cand in (f"{FONTS}/{name}", f"{FONTS}/Supplemental/{name}", f"{FONTS}/Helvetica.ttc"):
        try:
            return ImageFont.truetype(cand, size)
        except OSError:
            continue
    return ImageFont.load_default()


# ------------------------------------------------------------------------- 1. discharge summary (PDF)
def discharge_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(OUT / "Riverside_Discharge_Summary.pdf"), pagesize=letter, topMargin=40)
    P = lambda t, s="BodyText": Paragraph(t, st[s])
    tbl = lambda rows, w=None: Table(rows, colWidths=w, style=TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef3")),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [
        P("<b>RIVERSIDE MEDICAL CENTER</b> &nbsp; 1200 Riverside Pkwy, Springfield, IL 62701 &nbsp; Tel (217) 555-0100", "Normal"),
        P("HOSPITAL DISCHARGE SUMMARY", "Title"),
        tbl([["Patient", "VASQUEZ, Eleanor M.", "MRN", "RMC-0048213"],
             ["DOB", "06/22/1957 (69 y)", "Sex", "Female"],
             ["Admitted", "08/10/2026", "Discharged", "08/14/2026"],
             ["Attending", "Alan Brooks, MD (Hospitalist)", "PCP", "Priya Raman, MD"],
             ["Address", "42 Maple Ridge Dr, Springfield, IL 62704", "Phone (cell)", "(217) 555-0148"],
             ["Emergency contact", "Sofia Vasquez (daughter) (217) 555-0199", "Code status", "Full code"]],
            [80, 190, 80, 160]),
        Spacer(1, 10),
        P("<b>Principal diagnosis:</b> Acute on chronic heart failure with reduced ejection fraction (I50.23), EF 35% on TTE 08/11/2026."),
        P("<b>Secondary diagnoses:</b> Type 2 diabetes mellitus with hyperglycemia (E11.65); essential hypertension (I10); "
          "chronic kidney disease stage 3a (N18.31); hyperlipidemia (E78.5)."),
        P("<b>Hospital course:</b> 69-year-old woman presented to the ED on 08/10/2026 with 5 days of progressive dyspnea, "
          "orthopnea and 4 kg weight gain. BNP 890 pg/mL. Diuresed with IV furosemide, net -4.1 L. Transitioned to oral "
          "furosemide 40 mg daily. Echocardiogram showed EF 35% with global hypokinesis. Cardiology consult (Marcus Lee, MD) "
          "recommended outpatient follow-up for GDMT optimisation. Discharged home in stable condition with daughter."),
        P("<b>Procedures:</b> Transthoracic echocardiogram 08/11/2026 (CPT 93306), performed by Riverside Echo Lab."),
        Spacer(1, 6),
        P("<b>Allergies:</b> PENICILLIN - hives (moderate). No known food allergies."),
        P("<b>Discharge medications</b>"),
        tbl([["Medication", "Dose", "Route", "Frequency", "Comment"],
             ["Metformin (Glucophage)", "500 mg", "PO", "twice daily with meals", "continue"],
             ["Lisinopril", "10 mg", "PO", "once daily", "continue"],
             ["Furosemide (Lasix)", "40 mg", "PO", "once daily", "NEW"],
             ["Atorvastatin", "40 mg", "PO", "nightly", "continue"],
             ["Aspirin", "81 mg", "PO", "once daily", "continue"],
             ["Metoprolol succinate", "25 mg", "PO", "once daily", "NEW"]], [130, 50, 40, 120, 80]),
        Spacer(1, 8),
        P("<b>Key labs</b>"),
        tbl([["Test", "Result", "Units", "Ref range", "Flag", "Collected"],
             ["Hemoglobin A1c", "8.4", "%", "4.0-5.6", "H", "08/11/2026"],
             ["BNP", "890", "pg/mL", "<100", "H", "08/10/2026"],
             ["Creatinine", "1.3", "mg/dL", "0.5-1.1", "H", "08/13/2026"],
             ["eGFR", "52", "mL/min/1.73m2", ">60", "L", "08/13/2026"],
             ["Potassium", "4.4", "mmol/L", "3.5-5.1", "", "08/13/2026"]], [110, 50, 80, 70, 35, 75]),
        Spacer(1, 8),
        P("<b>Vitals at discharge (08/14/2026):</b> BP 128/76 mmHg, HR 72 bpm, SpO2 96% RA, weight 80.1 kg, height 162 cm."),
        P("<b>Social history:</b> Former smoker (quit 2009, 20 pack-years). Rare alcohol. Retired school teacher. Lives alone; daughter nearby."),
        P("<b>Immunizations:</b> Influenza vaccine given 08/13/2026 prior to discharge (lot FL2026-118)."),
        P("<b>Follow-up:</b> PCP Dr. Priya Raman within 7 days; Cardiology Dr. Marcus Lee in 4-6 weeks. Daily weights; call if gain > 1 kg/day."),
        Spacer(1, 10),
        P("Electronically signed: Alan Brooks, MD - 08/14/2026 16:42", "Italic"),
    ]
    doc.build(story)


# ------------------------------------------------------------------------- 2. PCP e-mail with CSV labs attachment
def pcp_email():
    labs = io.StringIO()
    w = csv.writer(labs)
    w.writerow(["Patient", "MRN", "Test", "Result", "Units", "Reference", "Flag", "Collected", "Lab"])
    for row in [("HbA1c", "7.9", "%", "4.0-5.6", "H"), ("Creatinine", "1.4", "mg/dL", "0.5-1.1", "H"),
                ("eGFR (CKD-EPI)", "48", "mL/min/1.73m2", ">60", "L"), ("Potassium", "5.1", "mmol/L", "3.5-5.1", "N"),
                ("LDL Cholesterol", "92", "mg/dL", "<100", "N")]:
        w.writerow(["Vasquez, Eleanor", "RMC-0048213", *row, "2026-09-01", "Quest Diagnostics"])

    msg = EmailMessage()
    msg["From"] = "Priya Raman MD <praman@springfieldfamilymed.example>"
    msg["To"] = "Marcus Lee MD <mlee@heartcareil.example>"
    msg["Cc"] = "referrals@heartcareil.example"
    msg["Date"] = "Thu, 03 Sep 2026 09:14:00 -0500"
    msg["Subject"] = "Referral: Eleanor Vasquez (DOB 6/22/57) - HFrEF follow-up, med changes"
    msg.set_content("""Hi Marcus,

Referring Mrs. Eleanor Vasquez (MRN RMC-0048213 at Riverside) for the post-discharge HF follow-up you
recommended during her August admission.

Saw her in clinic on 9/1. Doing better, weight down to 79 kg, BP 134/80, HR 70. A few changes since discharge:
 - Metformin increased from 500 mg to 1000 mg twice daily (A1c still 7.9).
 - Stopped lisinopril 10 mg; started sacubitril/valsartan (Entresto) 24/26 mg twice daily after 36 h washout.
 - Continues furosemide 40 mg daily, metoprolol succinate 25 mg daily, atorvastatin 40 mg, aspirin 81 mg.
 - She told us she also gets a rash with sulfa drugs (Bactrim years ago) - added to her allergy list.

Potassium is borderline at 5.1, so please keep an eye on it with the Entresto. Labs from 9/1 attached.

She has a new cell number: 217-555-0177 (the 0148 number is disconnected). Her daughter Sofia is
still her emergency contact. She fills everything at Walgreens #4471 on S Grand Ave.

Thanks,
Priya

Priya Raman, MD | Springfield Family Medicine | (217) 555-0123
""")
    msg.add_attachment(labs.getvalue().encode(), maintype="text", subtype="csv", filename="Vasquez_labs_2026-09-01.csv")
    (OUT / "PCP_referral_email.eml").write_bytes(bytes(msg))


# ------------------------------------------------------------------------- 3. insurance card photo
def insurance_card():
    from PIL import Image, ImageDraw, ImageFilter

    W, H = 1010, 640
    card = Image.new("RGB", (W, H), "#f4f8fc")
    d = ImageDraw.Draw(card)
    d.rounded_rectangle([0, 0, W - 1, H - 1], 36, fill="#f4f8fc", outline="#1c4f8c", width=6)
    d.rectangle([6, 6, W - 7, 120], fill="#1c4f8c")
    d.text((40, 30), "BlueCross BlueShield of Illinois", font=font("HelveticaNeue.ttc", 46), fill="white")
    d.text((W - 220, 40), "PPO", font=font("HelveticaNeue.ttc", 52), fill="#9fd0ff")
    rows = [("Subscriber", "ELEANOR M VASQUEZ"), ("Member ID", "XHB884512307"), ("Group No.", "0047731"),
            ("Plan", "Blue Choice PPO Silver"), ("DOB", "06/22/1957"), ("Effective", "01/01/2026"),
            ("RxBIN / PCN", "004336 / ADV"), ("PCP Copay $25   Specialist $50   ER $250", "")]
    y = 150
    for k, v in rows:
        d.text((50, y), k, font=font("HelveticaNeue.ttc", 28), fill="#4a5a6a")
        d.text((330, y), v, font=font("Courier.ttc", 34), fill="#0b1b2b")
        y += 56
    d.text((50, H - 60), "Customer Service 1-800-555-0110  |  bcbsil.example", font=font("HelveticaNeue.ttc", 24), fill="#4a5a6a")
    # photo-like: put on a desk background, rotate, blur slightly
    bg = Image.new("RGB", (1300, 950), "#8a7a66")
    rot = card.rotate(-4, expand=True, fillcolor="#8a7a66")
    bg.paste(rot, (120, 130))
    bg = bg.filter(ImageFilter.GaussianBlur(0.8))
    bg.save(OUT / "insurance_card_photo.jpg", quality=82)


# ------------------------------------------------------------------------- 4. patient intake form (docx)
def intake_docx():
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("HeartCare Illinois - New Patient Intake Form", 1)
    doc.add_paragraph("Clinic visit date: 9/18/2026        Provider: Dr. Marcus Lee (Cardiology)")
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in [("Full name", "Eleanor Maria Vasquez"), ("Date of birth", "06/23/1957"), ("Sex", "F"),
                 ("Address", "1187 Lakeview Ave, Apt 3B, Springfield IL 62702 (moved in Sept!)"),
                 ("Cell phone", "217-555-0177"), ("Email", "eleanor.vasquez57@example.com"),
                 ("Preferred language", "Spanish (English OK)"), ("Marital status", "Widowed"),
                 ("Emergency contact", "Sofia Vasquez-Hart (daughter) 217-555-0199"),
                 ("Primary care doctor", "Dr. Raman"),
                 ("Insurance", "BCBS IL PPO, member # XHB884512307, group 0047731"),
                 ("Preferred pharmacy", "Walgreens #4471, 300 S Grand Ave W, Springfield"),
                 ("Advance directive", "Yes - POLST on file at Riverside, daughter is healthcare POA")]:
        r = t.add_row().cells
        r[0].text, r[1].text = k, v
    doc.add_heading("Allergies (drug - what happens)", 2)
    doc.add_paragraph("Penicillin - my lips swelled up\nSulfa (Bactrim) - rash")
    doc.add_heading("Medicines I take now", 2)
    m = doc.add_table(rows=1, cols=3)
    m.style = "Table Grid"
    m.rows[0].cells[0].text, m.rows[0].cells[1].text, m.rows[0].cells[2].text = "Medicine", "How much", "How often"
    for row in [("Metformin", "1000mg", "2x a day"), ("Entresto", "24/26", "2x a day"), ("Lasix (water pill)", "20 mg", "morning"),
                ("Metoprolol", "25", "1x day"), ("Atorvastatin", "40 mg", "bedtime"), ("Baby aspirin", "81", "daily"),
                ("Vitamin D3", "2000 IU", "daily")]:
        c = m.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v
    doc.add_heading("Medical history (check all)", 2)
    doc.add_paragraph("[x] Heart failure   [x] Diabetes   [x] High blood pressure   [x] Kidney disease   [ ] Stroke\n"
                      "[x] High cholesterol   [ ] Cancer   Other: cataract surgery right eye 2019")
    doc.add_heading("Lifestyle", 2)
    doc.add_paragraph("Tobacco: quit in 2009.  Alcohol: a glass of wine on holidays.  Exercise: walking 15 min.\n"
                      "Occupation: retired teacher.  Lives: alone, daughter visits daily.")
    doc.add_paragraph("Weight today (self-reported): 172 lb")
    doc.add_paragraph("Signature: E. Vasquez        Date: 09/18/2026")
    doc.save(OUT / "HeartCare_intake_form.docx")


# ------------------------------------------------------------------------- 5. cardiology labs/vitals workbook
def labs_xlsx():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Labs"
    ws.append(["HeartCare Illinois - point of care + send-out results", None, None, None, None])
    ws.append(["Patient: VASQUEZ, ELEANOR  MRN RMC-0048213  DOB 1957-06-22", None, None, None, None])
    ws.append([])
    ws.append(["Analyte", "Value", "Unit", "Range", "Flag", "Collected"])
    for r in [("NT-proBNP", 1450, "pg/mL", "<125", "H"), ("Potassium", 5.3, "mmol/L", "3.5-5.1", "H"),
              ("Creatinine", 1.4, "mg/dL", "0.5-1.1", "H"), ("eGFR", 47, "mL/min/1.73m2", ">60", "L"),
              ("Sodium", 138, "mmol/L", "135-145", "")]:
        ws.append([*r, "2026-09-18"])
    v = wb.create_sheet("Vitals")
    v.append(["Date", "BP", "HR", "SpO2", "Weight (kg)", "Height (cm)", "BMI"])
    v.append(["2026-09-18", "138/82", 76, "97%", 78.2, 162, 29.8])
    wb.save(OUT / "HeartCare_labs_vitals_2026-09-18.xlsx")


# ------------------------------------------------------------------------- 6. HL7 ADT
def hl7():
    segs = [
        "MSH|^~\\&|HCI_REG|HEARTCARE_IL|RMC_EMPI|RIVERSIDE|20260918101522||ADT^A08|MSG00098231|P|2.5.1",
        "EVN|A08|20260918101522",
        "PID|1||RMC-0048213^^^RMC^MR~XHB884512307^^^BCBSIL^MB||VASQUEZ^ELEANOR^MARIA||19570622|F|||1187 LAKEVIEW AVE^APT 3B^SPRINGFIELD^IL^62702^USA||^PRN^CP^^^217^5550177|||W||||||H^Hispanic or Latino",
        "NK1|1|VASQUEZ-HART^SOFIA|DAU^Daughter|||^PRN^PH^^^217^5550199",
        "PV1|1|O|CARD^EXAM3^^HEARTCARE_IL||||1234567^LEE^MARCUS^^^MD|||CAR||||||||||||||||||||||||||||||||20260918",
        "AL1|1|DA|^PENICILLIN|MO|ANGIOEDEMA (LIP SWELLING)",
        "AL1|2|DA|^SULFONAMIDES|MI|RASH",
        "DG1|1||I50.22^Chronic systolic (congestive) heart failure^I10|||W",
        "IN1|1|BCBSPPO|00621|BLUECROSS BLUESHIELD OF ILLINOIS|||||0047731||||20260101|||VASQUEZ^ELEANOR^M|SEL|19570622|||||||||||||||||XHB884512307",
    ]
    (OUT / "ADT_A08_registration.hl7").write_text("\r".join(segs) + "\r")


# ------------------------------------------------------------------------- 7. scanned prescription
def prescription_scan():
    from PIL import Image, ImageDraw, ImageFilter

    W, H = 1100, 1400
    img = Image.new("RGB", (W, H), "#fbfaf4")
    d = ImageDraw.Draw(img)
    d.text((60, 50), "HeartCare Illinois", font=font("HelveticaNeue.ttc", 48), fill="#203040")
    d.text((60, 110), "Marcus Lee, MD - Cardiology   NPI 1234567893   DEA ML9876543", font=font("HelveticaNeue.ttc", 24), fill="#405060")
    d.text((60, 145), "800 E Capitol Ave, Springfield IL 62701  |  (217) 555-0160", font=font("HelveticaNeue.ttc", 24), fill="#405060")
    d.line([60, 190, W - 60, 190], fill="#203040", width=3)
    hw = font("Noteworthy.ttc", 40)
    lines = [("Patient:", "Eleanor Vasquez     DOB 6/22/57"), ("Date:", "9/18/2026"), ("", ""),
             ("Rx", "Furosemide 20 mg tab"), ("", "Sig: 1 tab PO every morning"), ("", "(DECREASE from 40 mg - K+ 5.3, wt down)"),
             ("", "Disp #30   Refills 3"), ("", ""), ("Rx", "Dapagliflozin (Farxiga) 10 mg tab"), ("", "Sig: 1 tab PO daily for HFrEF/CKD"),
             ("", "Disp #30   Refills 5"), ("", ""), ("", "Hold spironolactone for now."), ("", "Pharmacy: Walgreens #4471"),
             ("", ""), ("", "Recheck BMP in 1 week")]
    y = 240
    for k, v in lines:
        if k:
            d.text((60, y), k, font=font("HelveticaNeue.ttc", 30), fill="#203040")
        d.text((220, y), v, font=hw, fill="#1a2a8a")
        y += 62
    d.text((620, H - 160), "M. Lee MD", font=font("Noteworthy.ttc", 54), fill="#1a2a8a")
    d.line([600, H - 95, W - 80, H - 95], fill="#203040", width=2)
    # scanner look: noise, slight rotation, grey
    rnd = random.Random(7)
    px = img.load()
    for _ in range(25000):
        x, yy = rnd.randrange(W), rnd.randrange(H)
        g = rnd.randrange(150, 230)
        px[x, yy] = (g, g, g)
    img = img.rotate(1.6, expand=True, fillcolor="#e6e6e0").filter(ImageFilter.GaussianBlur(0.6)).convert("L")
    img.save(OUT / "Rx_scan_LeeMD_2026-09-18.png")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in (discharge_pdf, pcp_email, insurance_card, intake_docx, labs_xlsx, hl7, prescription_scan):
        fn()
        print("ok", fn.__name__)
    print("written to", OUT)


if __name__ == "__main__":
    main()

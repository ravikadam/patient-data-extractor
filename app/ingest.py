"""Turn an uploaded file into LLM content parts (+ plain text where we can get it).

PDFs and images go to the model natively (vision), so scanned documents and photos of cards work.
Office/e-mail/text formats are converted to text locally. E-mail attachments are split out as
their own documents so each one gets its own provenance.
"""
from __future__ import annotations

import base64
import csv
import email
import email.policy
import html
import io
import mimetypes
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}
CONVERTIBLE_IMAGES = {".tif", ".tiff", ".bmp"}
TEXT_TYPES = {".txt", ".md", ".hl7", ".json", ".xml", ".csv", ".tsv", ".log", ".rtf", ".ccd", ".fhir"}
SUPPORTED = set(IMAGE_TYPES) | CONVERTIBLE_IMAGES | TEXT_TYPES | {".pdf", ".docx", ".eml", ".html", ".htm", ".xlsx", ".xlsm"}

MAX_IMAGE_EDGE = 2000
MAX_TEXT_CHARS = 400_000


@dataclass
class Prepared:
    parts: list[dict[str, Any]]
    text: str | None = None                      # plain-text rendering, when available
    attachments: list[tuple[str, bytes]] = field(default_factory=list)  # (filename, bytes) to ingest separately
    hint_date: str | None = None                 # e.g. e-mail Date header


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


def _image_part(data: bytes, ext: str) -> dict[str, Any]:
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    needs_convert = ext in CONVERTIBLE_IMAGES or max(img.size) > MAX_IMAGE_EDGE or len(data) > 4_500_000
    if not needs_convert:
        return {"type": "image", "media_type": IMAGE_TYPES[ext], "data": _b64(data)}
    img.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="JPEG", quality=88)
    return {"type": "image", "media_type": "image/jpeg", "data": _b64(buf.getvalue())}


def _html_to_text(s: str) -> str:
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</h\d>|</li>", "\n", s)
    s = re.sub(r"(?i)</t[dh]>", " | ", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s)
    return re.sub(r"\n\s*\n+", "\n\n", s).strip()


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out: list[str] = []
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for ti, table in enumerate(d.tables):
        out.append(f"\n[Table {ti + 1}]")
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip()
                if not cells or cells[-1] != t:  # merged cells repeat
                    cells.append(t)
            out.append(" | ".join(cells))
    return "\n".join(out)


def _xlsx_text(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
    except Exception:
        return ""


def _eml(data: bytes) -> Prepared:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    headers = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h])
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        body = body_part.get_content()
        if body_part.get_content_type() == "text/html":
            body = _html_to_text(body)
    attachments = []
    for part in msg.iter_attachments():
        name = part.get_filename() or f"attachment{mimetypes.guess_extension(part.get_content_type()) or '.bin'}"
        payload = part.get_payload(decode=True)
        if payload:
            attachments.append((name, payload))
    hint = None
    if msg["Date"]:
        try:
            hint = email.utils.parsedate_to_datetime(msg["Date"]).date().isoformat()
        except Exception:
            pass
    att_note = ("\n\nAttachments (processed as separate documents): " + ", ".join(a[0] for a in attachments)) if attachments else ""
    text = f"{headers}\n\n{body}{att_note}"
    return Prepared(parts=[{"type": "text", "text": text}], text=text, attachments=attachments, hint_date=hint)


def prepare(filename: str, data: bytes) -> Prepared:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        # Best effort: treat as UTF-8 text if it decodes, otherwise reject.
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise ValueError(f"Unsupported file type: {ext or filename}")
        return Prepared(parts=[{"type": "text", "text": text[:MAX_TEXT_CHARS]}], text=text)

    if ext == ".pdf":
        text = _pdf_text(data)
        return Prepared(parts=[{"type": "pdf", "data": _b64(data), "name": filename}], text=text or None)
    if ext in IMAGE_TYPES or ext in CONVERTIBLE_IMAGES:
        return Prepared(parts=[_image_part(data, ext)])
    if ext == ".eml":
        return _eml(data)
    if ext == ".docx":
        text = _docx_text(data)
    elif ext in (".xlsx", ".xlsm"):
        text = _xlsx_text(data)
    elif ext in (".html", ".htm"):
        text = _html_to_text(data.decode("utf-8", errors="replace"))
    elif ext in (".csv", ".tsv"):
        raw = data.decode("utf-8", errors="replace")
        rows = csv.reader(io.StringIO(raw), delimiter="\t" if ext == ".tsv" else ",")
        text = "\n".join(" | ".join(r) for r in rows)
    else:
        text = data.decode("utf-8", errors="replace")
        if ext == ".hl7":
            text = text.replace("\r", "\n")
    text = text[:MAX_TEXT_CHARS]
    return Prepared(parts=[{"type": "text", "text": text}], text=text)
